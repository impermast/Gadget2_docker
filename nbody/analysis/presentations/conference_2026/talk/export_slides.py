#!/usr/bin/env python3
"""Build a faithful, offline PPTX from the current evidence-aware slides.md."""
import argparse, re, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

IMG = re.compile(r"!\[([^]]*)\]\(([^)]+)\)")

def clean(s):
    s = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", s)
    return re.sub(r"[*`]", "", s).strip()

def parse(path):
    slides, cur = [], None
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## ") and not line[3:].strip().startswith("Backup"):
            if cur: slides.append(cur)
            cur = {"title": clean(line[3:]), "lines": [], "images": []}
        elif line.startswith("### B") and re.match(r"### B0[1-5]\b", line):
            if cur: slides.append(cur)
            cur = {"title": clean(line[4:]), "lines": [], "images": []}
        elif cur is not None and line.strip():
            m = IMG.search(line)
            if m: cur["images"].append((m.group(1), m.group(2)))
            cur["lines"].append(clean(line))
    if cur: slides.append(cur)
    if len(slides) != 16:
        raise ValueError(f"expected 16 slides (S01-S11 + B01-B05), found {len(slides)}")
    return slides

def add_text(slide, text, x, y, w, h, size=16, color=(235,240,248), bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.name = "Calibri"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor(*color)

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--slides", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--asset-root", type=Path, help="base for relative figure paths (default: slides parent)")
    ap.add_argument("--allow-missing", action="store_true")
    args = ap.parse_args(argv)
    if not args.slides.is_file(): ap.error("slides file not found: " + str(args.slides))
    try: slides = parse(args.slides)
    except ValueError as e: ap.error(str(e))
    if not slides: ap.error("no ## slide headings found")
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for item in slides:
        s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(11,31,58)
        add_text(s, item["title"], .45, .25, 12.4, .7, 26, (79,195,247), True)
        y = 1.15
        for alt, rel in item["images"]:
            p = ((args.asset_root or args.slides.parent) / rel).resolve()
            if not p.is_file():
                if not args.allow_missing: ap.error(f"missing figure: {p}")
                print(f"warning: missing figure: {p}", file=sys.stderr); continue
        for i, (alt, rel) in enumerate(item["images"][:2]):
            p = ((args.asset_root or args.slides.parent) / rel).resolve()
            if not p.is_file(): continue
            left = .55 + i * 6.35
            s.shapes.add_picture(str(p), Inches(left), Inches(1.25), width=Inches(5.95), height=Inches(3.55))
        text_lines = [x for x in item["lines"] if not x.startswith("![") and not x.startswith("Source:")]
        tx, tw, ty = (.7, 12, 1.25)
        add_text(s, "\n".join("• " + x for x in text_lines), tx, ty if not item["images"] else 5.0, tw, 1.95 if item["images"] else 5.8, 12)
    args.output.parent.mkdir(parents=True, exist_ok=True); prs.save(args.output)
    print(f"created {args.output} slides={len(slides)} assets={sum(len(x['images']) for x in slides)}")

if __name__ == "__main__": raise SystemExit(main())
